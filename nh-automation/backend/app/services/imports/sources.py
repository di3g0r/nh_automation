"""Import sources: where catalog rows come from.

The real product/inventory data lives in an external database whose structure
is not known yet. Everything downstream (validation, preview, confirm) only
sees `SourceData` -- column names plus raw rows -- so the source can change
without touching the import rules:

- `ExternalDbSource`: any SQLAlchemy URL (`EXTERNAL_CATALOG_DB_URL`) plus a
  read-only SELECT per catalog from a .sql file. Usually enough: aliases,
  joins, CASE expressions and unit conversions all go in the SQL.
- `FileSource`: CSV / XLSX upload (manual fallback, and test fixture).
- Anything else (an API, a nightly export...): implement `ImportSource` and
  register it in `_registry()`.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

from sqlalchemy import create_engine, text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.pool import NullPool

from app.core.config import get_settings
from app.core.errors import AppError
from app.services.imports.columns import ImportKind

MAX_ROWS = 20_000
MAX_FILE_BYTES = 10 * 1024 * 1024
BACKEND_ROOT = Path(__file__).resolve().parents[3]


@dataclass
class SourceData:
    columns: list[str]
    # (row number shown to the user, {source column: raw value})
    rows: list[tuple[int, dict[str, Any]]] = field(default_factory=list)


@dataclass
class SourceInput:
    """What the request carries; each source picks what it needs."""

    filename: str | None = None
    content: bytes | None = None


class ImportSource(Protocol):
    name: str
    label: str  # Spanish, shown in the import dialog

    def is_available(self, kind: ImportKind) -> bool: ...

    def read(self, kind: ImportKind, data: SourceInput) -> SourceData: ...


def _too_many_rows() -> AppError:
    return AppError(
        "IMPORT_TOO_MANY_ROWS",
        f"La importación excede el máximo de {MAX_ROWS:,} filas.",
        422,
    )


# File source -----------------------------------------------------------------


class FileSource:
    name = "file"
    label = "Archivo CSV / XLSX"

    def is_available(self, kind: ImportKind) -> bool:
        return True

    def read(self, kind: ImportKind, data: SourceInput) -> SourceData:
        if not data.content or not data.filename:
            raise AppError("IMPORT_FILE_REQUIRED", "Seleccione un archivo CSV o XLSX.", 422)
        if len(data.content) > MAX_FILE_BYTES:
            raise AppError("IMPORT_FILE_TOO_LARGE", "El archivo excede 10 MB.", 422)
        name = data.filename.lower()
        if name.endswith(".xlsx"):
            return _read_xlsx(data.content)
        if name.endswith(".csv") or name.endswith(".txt"):
            return _read_csv(data.content)
        raise AppError(
            "IMPORT_FILE_TYPE", "Formato no soportado. Use un archivo .csv o .xlsx.", 422
        )


def _is_blank(values: list[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in values)


def _rows_from_matrix(matrix: list[tuple[int, list[Any]]]) -> SourceData:
    """First non-blank row is the header; blank rows are skipped."""
    header: list[str] | None = None
    out = SourceData(columns=[])
    for number, values in matrix:
        if _is_blank(values):
            continue
        if header is None:
            header = [str(v).strip() if v is not None else "" for v in values]
            out.columns = [h for h in header if h]
            continue
        row = {h: values[i] if i < len(values) else None for i, h in enumerate(header) if h}
        out.rows.append((number, row))
        if len(out.rows) > MAX_ROWS:
            raise _too_many_rows()
    if header is None:
        raise AppError("IMPORT_EMPTY", "El archivo no contiene datos.", 422)
    return out


def _read_csv(content: bytes) -> SourceData:
    # Excel in Spanish locales often saves CSV as cp1252 and with ';'.
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise AppError("IMPORT_ENCODING", "No se pudo leer el archivo (codificación).", 422)

    sample = decoded[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(decoded), dialect)
    return _rows_from_matrix([(i, list(r)) for i, r in enumerate(reader, start=1)])


def _read_xlsx(content: bytes) -> SourceData:
    from openpyxl import load_workbook  # local import: only needed here

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise AppError("IMPORT_FILE_INVALID", "El archivo XLSX no es válido.", 422) from exc
    try:
        ws = wb.worksheets[0]
        matrix = [
            (i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)
        ]
    finally:
        wb.close()
    return _rows_from_matrix(matrix)


# External database source ------------------------------------------------------

_READ_ONLY = re.compile(r"^\s*(select|with)\b", re.IGNORECASE)
_SQL_COMMENT = re.compile(r"(--[^\n]*\n)|(/\*.*?\*/)", re.DOTALL)


class ExternalDbSource:
    """Runs an admin-configured SELECT against the external catalog database.

    The query is configured by IT on the server (never from the UI) and its
    result columns go through the same alias mapping as file headers. The
    connection is opened per import and the transaction is always rolled back.
    TODO(external-db): once the real database is known, add its driver to
    requirements.txt and write external_sources/*.sql (see README there).
    """

    name = "external_db"
    label = "Base de datos externa"

    def __init__(self, url: str | None = None, query_files: dict[str, str] | None = None):
        settings = get_settings()
        self.url = url if url is not None else settings.external_catalog_db_url
        self.query_files = query_files or {
            "products": settings.external_products_query_file,
            "packaging-items": settings.external_packaging_query_file,
        }

    def _query_path(self, kind: ImportKind) -> Path:
        path = Path(self.query_files[kind])
        return path if path.is_absolute() else BACKEND_ROOT / path

    def is_available(self, kind: ImportKind) -> bool:
        return bool(self.url) and self._query_path(kind).is_file()

    def _load_query(self, kind: ImportKind) -> str:
        if not self.url:
            raise AppError(
                "EXTERNAL_SOURCE_NOT_CONFIGURED",
                "La base de datos externa no está configurada en el servidor.",
                422,
            )
        path = self._query_path(kind)
        if not path.is_file():
            raise AppError(
                "EXTERNAL_SOURCE_NOT_CONFIGURED",
                "Falta la consulta SQL para este catálogo en el servidor.",
                422,
                {"query_file": str(self.query_files[kind])},
            )
        sql = path.read_text(encoding="utf-8").strip().rstrip(";")
        if not _READ_ONLY.match(_SQL_COMMENT.sub("\n", sql)):
            raise AppError(
                "EXTERNAL_SOURCE_INVALID_QUERY",
                "La consulta de la base de datos externa debe ser un SELECT.",
                422,
            )
        return sql

    def read(self, kind: ImportKind, data: SourceInput) -> SourceData:
        sql = self._load_query(kind)
        engine = create_engine(self.url, poolclass=NullPool)
        try:
            with engine.connect() as conn:
                try:
                    result = conn.execute(text(sql))
                    columns = list(result.keys())
                    fetched = result.fetchmany(MAX_ROWS + 1)
                finally:
                    conn.rollback()
        except SQLAlchemyError as exc:
            raise AppError(
                "EXTERNAL_SOURCE_ERROR",
                "No se pudo leer la base de datos externa.",
                502,
                {"error": str(getattr(exc, "orig", None) or exc)[:500]},
            ) from exc
        finally:
            engine.dispose()

        if len(fetched) > MAX_ROWS:
            raise _too_many_rows()
        rows = [(i, dict(zip(columns, r, strict=True))) for i, r in enumerate(fetched, start=1)]
        return SourceData(columns=columns, rows=rows)


# Registry ----------------------------------------------------------------------


def _registry() -> dict[str, ImportSource]:
    # Built per call so configuration changes (and test overrides) apply.
    return {s.name: s for s in (FileSource(), ExternalDbSource())}


def get_source(name: str) -> ImportSource:
    source = _registry().get(name)
    if source is None:
        raise AppError("IMPORT_UNKNOWN_SOURCE", "Origen de importación desconocido.", 422)
    return source


def list_sources(kind: ImportKind) -> list[dict[str, Any]]:
    return [
        {"name": s.name, "label": s.label, "available": s.is_available(kind)}
        for s in _registry().values()
    ]
